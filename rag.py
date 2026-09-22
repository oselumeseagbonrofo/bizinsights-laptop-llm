"""
Offline RAG store for BizInsights LLM.

Zero new services. SQLite + FTS5 for candidate recall, pure-Python BM25
for reranking. The chat model stays the only ML process on the laptop.
"""
import math
import re
import secrets
import sqlite3
import threading
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

ROOT_DIR = Path(__file__).resolve().parent
DEFAULT_DB_PATH = ROOT_DIR / "rag_store.db"
DEFAULT_KNOWLEDGE_DIR = ROOT_DIR / "knowledge"

CHUNK_SIZE = 750
CHUNK_OVERLAP = 120
MAX_CONTEXT_CHARS = 2400

# Ingestion caps: a 10 MB upload must never explode into unbounded RAM/DB.
MAX_PDF_PAGES = 50
MAX_DOCX_PARAGRAPHS = 5000
MAX_DOCX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_TEXT_CHARS = 200 * 1024
MAX_CHUNKS_PER_DOC = 200

_WORD_RE = re.compile(r"[^\W_]+", re.UNICODE)
_CJK_RE = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]+")


import unicodedata


def _norm(text: str) -> str:
    """NFKC normalize so NFD and NFC forms of the same word match."""
    return unicodedata.normalize("NFKC", text or "")


def tokenize(text: str) -> List[str]:
    """Lowercase word tokens (unicode-aware) for BM25 and FTS5 query building."""
    return _WORD_RE.findall(_norm(text).casefold())


def chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> List[str]:
    """Split text into overlapping char windows, preferring word boundaries."""
    if not 0 <= overlap < size:
        raise ValueError("overlap must be smaller than chunk size")
    cleaned = re.sub(r"\r\n?", "\n", text or "")
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()
    if not cleaned:
        return []
    if len(cleaned) <= size:
        return [cleaned]
    chunks: List[str] = []
    start = 0
    while start < len(cleaned):
        end = min(start + size, len(cleaned))
        if end < len(cleaned):
            boundary = _boundary(cleaned, start, end)
            if boundary > start + size // 2:
                end = boundary
        piece = cleaned[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= len(cleaned):
            break
        start = max(end - overlap, start + 1)
    return chunks


def _boundary(text: str, start: int, end: int) -> int:
    """Last whitespace (space, tab, newline, NBSP) in [start, end)."""
    return max(text.rfind(" ", start, end), text.rfind("\n", start, end),
               text.rfind("\t", start, end), text.rfind("\xa0", start, end))


def load_file(path: Path) -> str:
    """Extract plain text from a supported upload. Raises ValueError if empty.

    All parser failures surface as ValueError so callers clean up the
    staged file and return 400 instead of 500.
    """
    suffix = path.suffix.lower()
    if suffix in (".txt", ".md", ".csv"):
        try:
            try:
                text = path.read_text(encoding="utf-8", errors="strict")
            except UnicodeDecodeError:
                # French shop docs often arrive in legacy Windows encoding.
                text = path.read_text(encoding="cp1252", errors="strict")
        except (UnicodeDecodeError, OSError) as e:
            raise ValueError(f"File is not readable text: {e}") from e
        if "\x00" in text:
            raise ValueError("File is not readable text.")
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise ValueError("PDF support needs 'pip install pypdf'.") from e
        try:
            reader = PdfReader(str(path))
            if len(reader.pages) > MAX_PDF_PAGES:
                raise ValueError(f"PDF has more than {MAX_PDF_PAGES} pages. Split it and retry.")
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"Could not read PDF: {e}") from e
    elif suffix == ".docx":
        try:
            from docx import Document
        except ImportError as e:
            raise ValueError("DOCX support needs 'pip install python-docx'.") from e
        try:
            _check_docx(str(path))
            doc = Document(str(path))
            parts = [p.text for p in doc.paragraphs]
            # Tables hold real content (contracts, price lists). Index them
            # instead of silently dropping what the user thinks is searchable.
            table_cells = 0
            for table in getattr(doc, "tables", []):
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text and cell.text.strip():
                            parts.append(cell.text.strip())
                            table_cells += 1
            if len(doc.paragraphs) + table_cells > MAX_DOCX_PARAGRAPHS:
                raise ValueError(f"DOCX has more than {MAX_DOCX_PARAGRAPHS} paragraphs. Split it and retry.")
            text = "\n".join(parts)
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"Could not read DOCX: {e}") from e
    else:
        raise ValueError(f"Unsupported file type: {suffix}")
    if not text.strip():
        raise ValueError("No readable text found in file (scanned PDFs have no text layer).")
    if len(text) > MAX_TEXT_CHARS:
        raise ValueError(f"Extracted text over {MAX_TEXT_CHARS // 1024}k chars. Split the file and retry.")
    return text


def _check_docx(path: str):
    """Reject zip bombs, macros, and entity-expansion payloads before parsing."""
    import zipfile
    try:
        with zipfile.ZipFile(path) as z:
            total = sum(i.file_size for i in z.infolist())
            if total > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise ValueError("DOCX expands over 50 MB. Split it and retry.")
            names = {i.filename for i in z.infolist()}
            # Macro carriers ride through paragraph-only scans and wait on
            # disk for a Windows admin to open them. Reject, do not store.
            if "word/vbaProject.bin" in names or "word/vbaData.xml" in names:
                raise ValueError("DOCX with macros is not accepted.")
            # Scan every document part, not just the first bytes of one:
            # entities can hide after padding, in headers and styles,
            # in customXml parts, or in relationship (.rels) parts.
            for item in z.infolist():
                name = item.filename
                if not (name.endswith(".xml") or name.endswith(".rels")):
                    continue
                if not (name.startswith("word/") or name.startswith("customXml/")):
                    continue
                if item.file_size > 2 * 1024 * 1024:
                    raise ValueError("DOCX part over 2 MB. Split it and retry.")
                with z.open(name) as fh:
                    blob = fh.read()
                if b"<!DOCTYPE" in blob or b"<!ENTITY" in blob:
                    raise ValueError("DOCX with embedded entities is not accepted.")
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Could not read DOCX: {e}") from e


def sanitize_fts_query(query: str, max_terms: int = 50) -> Optional[str]:
    """Build a safe FTS5 MATCH string: word tokens joined with OR.

    Tokens come from the unicode-aware tokenizer, so only word chars
    can reach the query (no quotes, parens, or operators survive).
    Single ASCII chars are dropped as noise; non-ASCII tokens are kept
    at length 1 since one CJK character carries meaning.
    Long queries keep their discriminative tail: beyond max_terms the
    longest terms win instead of the first terms.
    """
    terms = []
    for t in tokenize(query):
        try:
            t.encode("ascii")
            if len(t) >= 2:
                terms.append(t)
        except UnicodeEncodeError:
            terms.append(t)
    if len(terms) > max_terms:
        # Rare, long terms discriminate; filler heads do not. Deterministic
        # order (length, then alpha) so results do not wander per restart.
        # The shortest two terms keep a slot so a short code tail survives.
        by_len = sorted(set(terms), key=lambda t: (-len(t), t))
        shorts = sorted(set(terms), key=lambda t: (len(t), t))[:2]
        terms = by_len[: max(max_terms - len(shorts), 1)]
        for s in shorts:
            if s not in terms:
                terms.append(s)
    else:
        # Dedup preserving order for the common short-query case.
        terms = list(dict.fromkeys(terms))
        if len(terms) > max_terms:
            terms = terms[:max_terms]
    if not terms:
        return None
    # Quote bare FTS5 operators so a query like "AND OR" cannot break parsing.
    quoted = [f'"{t}"' if t in ("and", "or", "not", "near") else t for t in terms]
    return " OR ".join(quoted)


def _like_escape(term: str) -> str:
    """Escape a literal for a parameterized LIKE ... ESCAPE '\\' clause."""
    return term.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")


def cjk_terms(query: str) -> List[str]:
    """Contiguous CJK runs in the query. FTS porter cannot split these
    into searchable pieces, so they are matched with LIKE instead."""
    return _CJK_RE.findall(query or "")


def _stem(word: str) -> str:
    """Porter stemmer (public-domain algorithm) for ASCII words.

    Non-ASCII words pass through untouched. Replaces the earlier
    hand-rolled suffix rules, which produced wrong roots
    (running to runn, agreed to agre) and hid cross-form misses.
    """
    try:
        word.encode("ascii")
    except UnicodeEncodeError:
        return word
    w = word.lower()
    if len(w) <= 2:
        return w
    return _porter(w)


def _porter(w: str) -> str:
    vowels = "aeiou"

    def is_consonant(s, i):
        c = s[i]
        if c in vowels:
            return False
        if c == "y":
            return True if i == 0 else not is_consonant(s, i - 1)
        return True

    def measure(s):
        # Canonical Porter m: count VC sequences in the C/V pattern.
        # A trailing vowel group must not count (tree->0, trouble->1).
        cv = "".join("c" if is_consonant(s, i) else "v" for i in range(len(s)))
        return cv.count("vc")

    def has_vowel(s):
        return any(not is_consonant(s, i) for i in range(len(s)))

    def double_consonant(s):
        return len(s) >= 2 and is_consonant(s, -1) and is_consonant(s, -2) and s[-1] == s[-2]

    def cvc(s):
        if len(s) < 3 or not is_consonant(s, -1) or is_consonant(s, -2) or not is_consonant(s, -3):
            return False
        return s[-1] not in "wxy"

    def step1a(s):
        if s.endswith("sses"):
            return s[:-2]
        if s.endswith("ies"):
            return s[:-2]
        if s.endswith("ss"):
            return s
        if s.endswith("s"):
            return s[:-1]
        return s

    def step1b(s):
        flag = False
        if s.endswith("eed"):
            if measure(s[:-3]) > 0:
                return s[:-1]
            return s
        for suffix in ("ed", "ing"):
            if s.endswith(suffix) and has_vowel(s[: -len(suffix)]):
                s = s[: -len(suffix)]
                flag = True
                break
        if not flag:
            return s
        if s.endswith(("at", "bl", "iz")):
            return s + "e"
        if double_consonant(s) and s[-1] not in "lsz":
            return s[:-1]
        if measure(s) == 1 and cvc(s):
            return s + "e"
        return s

    def step1c(s):
        if s.endswith("y") and has_vowel(s[:-1]):
            return s[:-1] + "i"
        return s

    def step2(s):
        for suffix, rep, m in (
            ("ational", "ate", 0), ("tional", "tion", 0), ("enci", "ence", 0),
            ("anci", "ance", 0), ("izer", "ize", 0), ("bli", "ble", 0),
            ("alli", "al", 0), ("entli", "ent", 0), ("eli", "e", 0),
            ("ousli", "ous", 0), ("ization", "ize", 0), ("ation", "ate", 0),
            ("ator", "ate", 0), ("alism", "al", 0), ("iveness", "ive", 0),
            ("fulness", "ful", 0), ("ousness", "ous", 0), ("aliti", "al", 0),
            ("iviti", "ive", 0), ("biliti", "ble", 0), ("logi", "log", 0),
        ):
            if s.endswith(suffix):
                # First-match-wins: a matched suffix that fails its guard
                # must not fall through to a shorter overlapping suffix.
                if measure(s[: -len(suffix)]) > m:
                    return s[: -len(suffix)] + rep
                return s
        return s

    def step3(s):
        for suffix, rep, m in (
            ("icate", "ic", 0), ("ative", "", 0), ("alize", "al", 0),
            ("iciti", "ic", 0), ("ical", "ic", 0), ("ful", "", 0), ("ness", "", 0),
        ):
            if s.endswith(suffix) and measure(s[: -len(suffix)]) > m:
                return s[: -len(suffix)] + rep
        return s

    def step4(s):
        for suffix in ("al", "ance", "ence", "er", "ic", "able", "ible",
                       "ant", "ement", "ment", "ent", "ou", "ism", "ate",
                       "iti", "ous", "ive", "ize"):
            if s.endswith(suffix):
                # First-match-wins: ement matched must never try ment/ent.
                if measure(s[: -len(suffix)]) > 1:
                    return s[: -len(suffix)]
                return s
        if s.endswith("sion") or s.endswith("tion"):
            if measure(s[:-3]) > 1:
                return s[:-3]
            return s
        return s

    def step5(s):
        if s.endswith("e"):
            m = measure(s[:-1])
            if m > 1 or (m == 1 and not cvc(s[:-1])):
                s = s[:-1]
        if s.endswith("ll") and measure(s) > 1:
            s = s[:-1]
        return s

    if w.startswith("y"):
        w = "Y" + w[1:]
    for fn in (step1a, step1b, step1c, step2, step3, step4, step5):
        w = fn(w)
    return w.lower()


def bm25_scores(query_terms: List[str], docs_terms: List[List[str]]) -> List[float]:
    """Pure-Python BM25 (k1=1.5, b=0.75). Returns one score per doc.

    Query terms are deduplicated first: repeating a word three times
    must not triple its weight.
    """
    k1, b = 1.5, 0.75
    query_terms = list(dict.fromkeys(_stem(t) for t in query_terms))
    docs_terms = [[_stem(t) for t in terms] for terms in docs_terms]
    n = len(docs_terms)
    if n == 0 or not query_terms:
        return [0.0] * n
    df: Dict[str, int] = {}
    for terms in docs_terms:
        for t in set(terms):
            df[t] = df.get(t, 0) + 1
    avg_len = sum(len(t) for t in docs_terms) / max(n, 1)
    scores: List[float] = []
    for terms in docs_terms:
        tf: Dict[str, int] = {}
        for t in terms:
            tf[t] = tf.get(t, 0) + 1
        score = 0.0
        doc_len = len(terms) or 1
        for q in query_terms:
            f = tf.get(q, 0)
            if f == 0:
                continue
            idf = math.log((n - df.get(q, 0) + 0.5) / (df.get(q, 0) + 0.5) + 1.0)
            denom = f + k1 * (1 - b + b * doc_len / max(avg_len, 1))
            score += idf * f * (k1 + 1) / denom
        scores.append(score)
    return scores


class RAGStore:
    def __init__(self, db_path: Path = DEFAULT_DB_PATH):
        self.db_path = Path(db_path)
        self._init_db()

    def _connect(self) -> sqlite3.Connection:
        conn = sqlite3.connect(str(self.db_path), timeout=10.0)
        conn.row_factory = sqlite3.Row
        try:
            conn.execute("PRAGMA journal_mode=WAL;")
        except sqlite3.OperationalError:
            pass
        try:
            conn.execute("PRAGMA busy_timeout=10000;")
        except sqlite3.OperationalError:
            pass
        return conn

    def _init_db(self):
        with self._connect() as conn:
            conn.execute("""
                CREATE TABLE IF NOT EXISTS documents (
                    doc_id TEXT PRIMARY KEY,
                    filename TEXT NOT NULL,
                    added_at TEXT NOT NULL,
                    num_chunks INTEGER NOT NULL
                );
            """)
            conn.execute("""
                CREATE TABLE IF NOT EXISTS chunks (
                    doc_id TEXT NOT NULL,
                    chunk_id INTEGER NOT NULL,
                    text TEXT NOT NULL,
                    PRIMARY KEY (doc_id, chunk_id)
                );
            """)
            try:
                conn.execute("""
                    CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts
                    USING fts5(doc_id, chunk_id UNINDEXED, text, tokenize='porter');
                """)
            except sqlite3.OperationalError:
                pass  # FTS5 unavailable; LIKE + BM25 fallback covers it.

    def _fts_ok(self, conn: sqlite3.Connection) -> bool:
        try:
            conn.execute("SELECT count(*) FROM chunks_fts LIMIT 1").fetchone()
            return True
        except sqlite3.OperationalError:
            return False

    def _rebuild_fts(self, conn: sqlite3.Connection):
        """Reconcile the FTS index with the chunks table after a partial failure."""
        conn.execute("DELETE FROM chunks_fts;")
        conn.execute("""
            INSERT INTO chunks_fts (doc_id, chunk_id, text)
            SELECT doc_id, chunk_id, text FROM chunks;
        """)

    def check_consistency(self) -> Dict[str, int]:
        """Row counts per table. chunks and chunks_fts must agree per doc."""
        with self._connect() as conn:
            docs = conn.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
            chunks = conn.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
            try:
                fts = conn.execute("SELECT COUNT(*) FROM chunks_fts").fetchone()[0]
            except sqlite3.OperationalError:
                fts = -1
            return {"documents": docs, "chunks": chunks, "chunks_fts": fts}

    def add_document(self, filename: str, text: str) -> Dict:
        # NFKC at ingest so NFD and NFC forms of the same word match at
        # query time. Tradeoff: compatibility chars (ligatures, fullwidth)
        # persist normalized, so displayed source text can differ from the
        # uploaded bytes. Recall wins over byte fidelity here.
        text = _norm(text)
        chunks = chunk_text(text)
        if not chunks:
            raise ValueError("No readable text found in file.")
        if len(chunks) > MAX_CHUNKS_PER_DOC:
            raise ValueError(f"Document splits into more than {MAX_CHUNKS_PER_DOC} chunks. Split it and retry.")
        doc_id = secrets.token_hex(8)
        added_at = datetime.utcnow().isoformat()
        with self._connect() as conn:
            try:
                conn.execute("BEGIN IMMEDIATE;")
            except sqlite3.OperationalError as e:
                # Never proceed on weaker isolation: contention must fail
                # fast as 429, not limp on deferred and die at commit as 400.
                try:
                    conn.rollback()
                except sqlite3.OperationalError:
                    pass
                raise sqlite3.OperationalError("database is locked") from e
            # Total FTS absence (no FTS5 in this SQLite build) degrades to
            # LIKE + BM25 instead of failing every upload. Partial FTS
            # failure below still rolls back, never half-indexes.
            fts_available = self._fts_ok(conn)
            try:
                conn.execute(
                    "INSERT INTO documents (doc_id, filename, added_at, num_chunks) VALUES (?, ?, ?, ?)",
                    (doc_id, filename, added_at, len(chunks)),
                )
                for idx, piece in enumerate(chunks):
                    conn.execute(
                        "INSERT INTO chunks (doc_id, chunk_id, text) VALUES (?, ?, ?)",
                        (doc_id, idx, piece),
                    )
                    # Fail closed: a chunk missing from FTS is a silent
                    # recall hole, so roll the whole document back instead.
                    if fts_available:
                        conn.execute(
                            "INSERT INTO chunks_fts (doc_id, chunk_id, text) VALUES (?, ?, ?)",
                            (doc_id, idx, piece),
                        )
            except sqlite3.OperationalError as e:
                conn.rollback()
                msg = str(e).lower()
                if "locked" in msg or "busy" in msg:
                    raise sqlite3.OperationalError("database is locked") from e
                raise ValueError(f"Search index write failed, document not stored: {e}") from e
        return {"doc_id": doc_id, "filename": filename, "num_chunks": len(chunks), "added_at": added_at}

    def delete_document(self, doc_id: str) -> bool:
        with self._connect() as conn:
            try:
                conn.execute("BEGIN IMMEDIATE;")
            except sqlite3.OperationalError as e:
                try:
                    conn.rollback()
                except sqlite3.OperationalError:
                    pass
                raise sqlite3.OperationalError("database is locked") from e
            cur = conn.execute("SELECT doc_id FROM documents WHERE doc_id = ?", (doc_id,))
            if not cur.fetchone():
                conn.rollback()
                return False
            fts_available = self._fts_ok(conn)
            conn.execute("DELETE FROM chunks WHERE doc_id = ?", (doc_id,))
            if fts_available:
                try:
                    conn.execute("DELETE FROM chunks_fts WHERE doc_id = ?", (doc_id,))
                except sqlite3.OperationalError as e:
                    # Contention must read as 429 like every other busy path,
                    # not as a 400 index failure.
                    if "locked" in str(e).lower() or "busy" in str(e).lower():
                        try:
                            conn.rollback()
                        except sqlite3.OperationalError:
                            pass
                        raise sqlite3.OperationalError("database is locked") from e
                    # FTS out of sync: rebuild it from the surviving chunks
                    # rather than leave deleted text searchable.
                    try:
                        self._rebuild_fts(conn)
                    except sqlite3.OperationalError as e2:
                        conn.rollback()
                        if "locked" in str(e2).lower() or "busy" in str(e2).lower():
                            raise sqlite3.OperationalError("database is locked") from e2
                        raise ValueError("Search index unavailable, delete aborted.") from e2
            conn.execute("DELETE FROM documents WHERE doc_id = ?", (doc_id,))
            return True

    def list_documents(self) -> List[Dict]:
        with self._connect() as conn:
            rows = conn.execute(
                "SELECT doc_id, filename, added_at, num_chunks FROM documents ORDER BY added_at DESC"
            ).fetchall()
            return [dict(r) for r in rows]

    def _doc_meta(self, conn: sqlite3.Connection) -> Dict[str, Dict]:
        meta: Dict[str, Dict] = {}
        for r in conn.execute("SELECT doc_id, filename, num_chunks FROM documents").fetchall():
            meta[r["doc_id"]] = {"filename": r["filename"], "num_chunks": r["num_chunks"]}
        return meta

    def search(self, query: str, top_k: int = 3) -> List[Dict]:
        query = (query or "").strip()
        if not query:
            return []
        with self._connect() as conn:
            meta = self._doc_meta(conn)
            if not meta:
                return []
            seen = set()
            candidates: List[Dict] = []

            def _add(doc_id, chunk_id, text):
                key = (doc_id, chunk_id)
                if key not in seen and doc_id in meta:
                    seen.add(key)
                    candidates.append({"doc_id": doc_id, "chunk_id": int(chunk_id), "text": text})

            # 1. FTS recall for word queries (overfetch; BM25 reranks).
            fts_q = sanitize_fts_query(query)
            if fts_q is not None:
                try:
                    rows = conn.execute(
                        "SELECT doc_id, chunk_id, text FROM chunks_fts "
                        "WHERE chunks_fts MATCH ? ORDER BY rank LIMIT 50",
                        (fts_q,),
                    ).fetchall()
                    for r in rows:
                        _add(r["doc_id"], r["chunk_id"], r["text"])
                except sqlite3.OperationalError:
                    pass

            # 2. CJK runs: FTS porter cannot split them, so match substrings.
            # Wider pool than before; BM25/coverage ranks after, not SQL recency.
            for run in cjk_terms(_norm(query))[:6]:
                try:
                    rows = conn.execute(
                        "SELECT doc_id, chunk_id, text FROM chunks "
                        "WHERE text LIKE ? ESCAPE '\\' ORDER BY rowid DESC LIMIT 50",
                        (f"%{_like_escape(run)}%",),
                    ).fetchall()
                    for r in rows:
                        _add(r["doc_id"], r["chunk_id"], r["text"])
                except sqlite3.OperationalError:
                    pass

            # 3. Fallback when recall found nothing: per-term LIKE in SQL.
            # Score-then-cut: take up to 200 into BM25, cut after ranking.
            # Terms are longest-first (deterministic) so a discriminative
            # tail survives; shortest terms keep slots for code IDs.
            if not candidates:
                raw = [t for t in tokenize(query) if len(t) >= 2]
                uniq = set(raw)
                by_len = sorted(uniq, key=lambda t: (-len(t), t))
                shorts = sorted(uniq, key=lambda t: (len(t), t))[:2]
                terms = by_len[:8]
                for s in shorts:
                    if s not in terms:
                        terms.append(s)
                terms = terms[:10] or raw[:10]
                if terms:
                    conds = " OR ".join(["text LIKE ? ESCAPE '\\'"] * len(terms))
                    try:
                        rows = conn.execute(
                            f"SELECT doc_id, chunk_id, text FROM chunks WHERE {conds} "
                            f"ORDER BY rowid DESC LIMIT 200",
                            tuple(f"%{_like_escape(t)}%" for t in terms),
                        ).fetchall()
                        for r in rows:
                            _add(r["doc_id"], r["chunk_id"], r["text"])
                    except sqlite3.OperationalError:
                        pass
            if not candidates:
                # 4. Last resort: the literal query as a substring. Gated at
                # 3 chars so "a i" style fragments keep returning empty.
                # Nonsense still returns empty; CJK and code IDs survive.
                # Try the full head plus the longest token so a tail unique
                # term is reachable even in a long filler query.
                needles = []
                head = _norm(query)[:60].casefold().strip()
                if len(head) >= 3:
                    needles.append(head)
                tails = sorted(set(tokenize(query)), key=lambda t: (-len(t), t))[:2]
                shorts = sorted(set(tokenize(query)), key=lambda t: (len(t), t))[:1]
                for t in tails + [s for s in shorts if s not in tails]:
                    if len(t) >= 2 and t not in head:
                        needles.append(t)
                for needle in needles[:4]:
                    try:
                        rows = conn.execute(
                            "SELECT doc_id, chunk_id, text FROM chunks "
                            "WHERE lower(text) LIKE ? ESCAPE '\\' ORDER BY rowid DESC LIMIT 30",
                            (f"%{_like_escape(needle)}%",),
                        ).fetchall()
                        for r in rows:
                            _add(r["doc_id"], r["chunk_id"], r["text"])
                    except sqlite3.OperationalError:
                        pass
                    if candidates:
                        break
            if not candidates:
                return []
            q_terms = tokenize(query)
            docs_terms = [tokenize(c["text"]) for c in candidates]
            scores = bm25_scores(q_terms, docs_terms)
            cjk_runs = cjk_terms(_norm(query))
            ranked = sorted(zip(candidates, scores), key=lambda p: p[1], reverse=True)
            scored: Dict[tuple, Dict] = {}
            for cand, score in ranked[: max(top_k, 1)]:
                if score <= 0:
                    continue
                m = meta.get(cand["doc_id"], {"filename": "unknown", "num_chunks": 0})
                scored[(cand["doc_id"], cand["chunk_id"])] = {
                    "doc_id": cand["doc_id"],
                    "filename": m["filename"],
                    "chunk_id": cand["chunk_id"],
                    "num_chunks": m["num_chunks"],
                    "text": cand["text"],
                    "score": round(float(score), 4),
                }
            if cjk_runs:
                # BM25 cannot match CJK substrings to whole-run tokens, so
                # score them by query coverage instead of a fixed constant:
                # a full-query CJK match outranks a weak Latin one, and a
                # mixed query returns both pools instead of either/or.
                total = max(sum(len(r) for r in cjk_runs), 1)
                for cand in candidates:
                    low = _norm(cand["text"]).casefold()
                    matched = sum(len(r) for r in cjk_runs if r.casefold() in low)
                    if not matched:
                        continue
                    key = (cand["doc_id"], cand["chunk_id"])
                    cover = round(min(0.5 + 0.5 * matched / total, 1.0), 4)
                    if key not in scored or cover > scored[key]["score"]:
                        m = meta.get(cand["doc_id"], {"filename": "unknown", "num_chunks": 0})
                        scored[key] = {
                            "doc_id": cand["doc_id"],
                            "filename": m["filename"],
                            "chunk_id": cand["chunk_id"],
                            "num_chunks": m["num_chunks"],
                            "text": cand["text"],
                            "score": cover,
                        }
            hits = sorted(scored.values(), key=lambda h: h["score"], reverse=True)
            return hits[: max(top_k, 1)]

    def format_context(self, hits: List[Dict], max_chars: int = MAX_CONTEXT_CHARS) -> str:
        parts: List[str] = []
        used = 0
        for h in hits:
            header = f"[source: {h['filename']} | chunk {h['chunk_id'] + 1}/{h['num_chunks']}]"
            body = h["text"].strip()
            block = f"{header}\n{body}"
            if used + len(block) > max_chars:
                # Charge the "..." marker inside the budget, never on top.
                room = max_chars - used - len(header) - 1 - 3
                if room > 100:
                    parts.append(f"{header}\n{body[:room].rstrip()}...")
                break
            parts.append(block)
            used += len(block) + 2
        return "\n\n".join(parts)


_STORE: Optional[RAGStore] = None
_STORE_LOCK = threading.Lock()


def get_store(db_path: Path = DEFAULT_DB_PATH) -> RAGStore:
    """Process-wide singleton for the default path; fresh instance otherwise."""
    global _STORE
    if str(db_path) == str(DEFAULT_DB_PATH):
        with _STORE_LOCK:
            if _STORE is None:
                _STORE = RAGStore(db_path)
            return _STORE
    return RAGStore(db_path)
